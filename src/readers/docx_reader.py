"""DOCX document reader using python-docx."""

from pathlib import Path

from docx import Document

from src.readers.base_reader import BaseReader, DocumentSegment
from src.utils.logger import setup_logger

logger = setup_logger(__name__)


class DOCXReader(BaseReader):
    """Extract text from DOCX files, preserving paragraph structure."""

    def read(self, file_path: Path) -> list[DocumentSegment]:
        segments: list[DocumentSegment] = []

        try:
            doc = Document(str(file_path))
            paragraphs: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    style_name = para.style.name if para.style else "Normal"
                    if style_name.startswith("Heading"):
                        paragraphs.append(f"\n{text}\n")
                    else:
                        paragraphs.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
            if full_text.strip():
                segments.append(
                    DocumentSegment(
                        text=full_text,
                        page_number=None,
                        metadata={"sections": len(paragraphs)},
                    )
                )
        except Exception as exc:
            logger.error("Failed to read DOCX %s: %s", file_path, exc)
            raise RuntimeError(f"Could not read DOCX file: {file_path}") from exc

        logger.info("DOCX reader extracted content from %s", file_path.name)
        return segments
